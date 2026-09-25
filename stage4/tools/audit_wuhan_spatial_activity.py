"""Stream published Figure 2/4 source tables; never execute upstream code.

Three five-column blocks are read independently, retaining missing vs zero.
Only compact aggregate results are written, never workbook edits or identifiers.
"""
from collections import Counter, defaultdict
from pathlib import Path
import hashlib
import json
import math
import time

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
PUBLIC = ROOT / 'stage4/output/external_research/wuhan_robotaxi_2026'
DOC = ROOT / 'stage4/docs/wuhan_robotaxi_review/spatial_activity'


def sha(path):
    with path.open('rb') as handle:
        return hashlib.file_digest(handle, 'sha256').hexdigest()


def overlap(grid_counts):
    a = {k for k, v in grid_counts.items() if v[0] > 0}
    h = {k for k, v in grid_counts.items() if v[1] > 0}
    total_a = sum(v[0] for v in grid_counts.values())
    total_h = sum(v[1] for v in grid_counts.values())
    common = a & h
    return {'represented_grids': len(grid_counts), 'robotaxi_active_grids': len(a),
            'hv_active_grids': len(h), 'both_active_grids': len(common),
            'either_active_grids': len(a | h),
            'active_grid_jaccard': len(common) / len(a | h) if a | h else None,
            'robotaxi_count': total_a, 'hv_count': total_h,
            'robotaxi_mass_in_shared_active_grids': sum(grid_counts[k][0] for k in common) / total_a if total_a else None,
            'hv_mass_in_shared_active_grids': sum(grid_counts[k][1] for k in common) / total_h if total_h else None,
            'normalized_spatial_distribution_overlap': sum(min(v[0]/total_a, v[1]/total_h) for v in grid_counts.values()) if total_a and total_h else None}


def audit_figure2(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book['source_data_figure2b']
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator)
    block_header = ('gx', 'gy', 'hour_bin', 'Robotaxi_Order_Count', 'HV_Order_Count')
    assert all(tuple(header[i:i+5]) == block_header for i in [0, 5, 10])
    block_info = [Counter() for _ in range(3)]
    controls = {}
    seen = {}  # one compact timestamp bitset per grid, not millions of tuples
    time_slots = {}
    temporal = defaultdict(lambda: Counter())
    grids = defaultdict(lambda: [0, 0])
    hours = defaultdict(lambda: defaultdict(lambda: [0, 0]))
    day = defaultdict(lambda: Counter())
    daily_grids = defaultdict(lambda: defaultdict(lambda: [0, 0]))
    states = Counter()
    duplicates = 0
    bounds = [math.inf, -math.inf, math.inf, -math.inf]
    for row_number, row in enumerate(iterator, 2):
        if row[15] is not None:
            assert row[15] not in controls
            controls[str(row[15])] = row[16]
        for b, offset in enumerate([0, 5, 10]):
            values = row[offset:offset+5]
            if all(v is None for v in values):
                block_info[b]['blank_padding_rows'] += 1; continue
            if any(v is None for v in values):
                block_info[b]['partial_missing_rows'] += 1; continue
            x, y, stamp, a, h = values
            assert all(isinstance(v, (int, float)) and float(v).is_integer() and v >= 0 for v in [x, y, a, h])
            x, y, a, h = map(int, [x, y, a, h])
            stamp = str(stamp)
            assert len(stamp) == 19 and stamp[13:] == ':00:00'
            hour = int(stamp[11:13]); assert 0 <= hour < 24
            date = stamp[:10]
            slot = time_slots.setdefault(stamp, len(time_slots))
            key = (x, y); bit = 1 << slot
            if seen.get(key, 0) & bit:
                duplicates += 1
                continue  # mark result unusable for definitive totals if duplicates occur
            seen[key] = seen.get(key, 0) | bit
            bounds = [min(bounds[0], x), max(bounds[1], x), min(bounds[2], y), max(bounds[3], y)]
            block_info[b]['records'] += 1
            block_info[b]['robotaxi_count'] += a; block_info[b]['hv_count'] += h
            state = 'both_positive' if a and h else 'robotaxi_only_positive' if a else 'hv_only_positive' if h else 'both_zero'
            states[state] += 1
            temporal[stamp]['cells'] += 1
            temporal[stamp]['robotaxi'] += a; temporal[stamp]['hv'] += h
            temporal[stamp][state] += 1
            day[date]['cells'] += 1; day[date]['robotaxi'] += a; day[date]['hv'] += h
            grids[key][0] += a; grids[key][1] += h
            daily_grids[date][key][0] += a; daily_grids[date][key][1] += h
            hours[hour][key][0] += a; hours[hour][key][1] += h
        if row_number % 200000 == 0:
            print('figure2b worksheet rows scanned:', row_number, flush=True)
    published_top30 = {}
    for row in book['source_data_figure2c_line'].iter_rows(min_row=2, values_only=True):
        if row[1] == 30: published_top30[str(row[0])] = row[2]
    radar = list(book['source_data_figure2c_radar'].iter_rows(values_only=True))
    book.close()
    assert duplicates == 0, f'duplicate keys: {duplicates}; do not silently aggregate'
    assert all(b['partial_missing_rows'] == 0 for b in block_info)
    total = overlap(grids)
    assert total['robotaxi_count'] == sum(v['robotaxi'] for v in temporal.values())
    assert total['hv_count'] == sum(v['hv'] for v in temporal.values())
    total_cells = sum(states.values())
    expected = len(grids) * len(time_slots)
    count_matches = {k: total[v] == controls[k] for k, v in [('Robotaxi','robotaxi_count'), ('HV','hv_count')]}
    hourly = []
    for hour, frame in sorted(hours.items()):
        row = {'hour': hour, **overlap(frame)}
        ts = [v for k, v in temporal.items() if int(k[11:13]) == hour]
        row['published_time_slots'] = len(ts)
        row['cell_states'] = {s: sum(v[s] for v in ts) for s in states}
        hourly.append(row)
    complete_mask = (1 << len(time_slots)) - 1
    shared_dates = [d for d, v in day.items() if v['robotaxi'] > 0 and v['hv'] > 0]
    shared_grids = defaultdict(lambda: [0, 0])
    for d in shared_dates:
        for key, counts in daily_grids[d].items():
            shared_grids[key][0] += counts[0]
            shared_grids[key][1] += counts[1]
    return {'sheet': sheet.title, 'declared_rows_including_header': sheet.max_row,
            'blocks': [dict(v) for v in block_info], 'published_total_controls': controls,
            'totals_match_published_controls': count_matches,
            'duplicate_grid_timestamp_keys': duplicates, 'unique_grid_count': len(grids),
            'grid_index_bounds': bounds, 'unique_published_timestamps': len(time_slots),
            'first_published_timestamp': min(time_slots), 'last_published_timestamp': max(time_slots),
            'date_count': len(day), 'represented_cell_count': total_cells,
            'rectangular_grid_timestamp_product': expected,
            'all_grids_have_all_published_timestamps': all(v == complete_mask for v in seen.values()),
            'cell_states': dict(states), 'spatial_totals': total, 'hourly': hourly,
            'both_modes_positive_total_dates': sorted(shared_dates),
            'both_modes_positive_date_spatial_totals': overlap(shared_grids),
            'both_modes_positive_date_cell_states': {s: sum(v[s] for k,v in temporal.items() if k[:10] in shared_dates) for s in states},
            'daily': [{'date': k, **dict(v), 'published_hours': sum(t.startswith(k) for t in time_slots)} for k,v in sorted(day.items())],
            'published_figure2c_top30_percent': published_top30,
            'published_figure2c_radar': radar,
            'timing_caveat': 'Upstream fig2b maps HV dates by sorted ordinal to AV dates; actual calendar synchrony not established',
            'zeros_caveat': 'Upstream grid x observed-hour expansion fills missing group counts with zero; zero does not establish geofence or capability'}


def audit_figure4(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book['figure4a'].iter_rows(min_row=2, values_only=True))
    published = {r[6]:r[7] for r in rows if r[6] is not None}
    widths = [r[2]-r[1] for r in rows]
    assert all(w > 0 for w in widths)
    assert all(abs(rows[i][2]-rows[i+1][1]) < 1e-10 for i in range(len(rows)-1))
    masses = {label: [r[col]*w for r,w in zip(rows,widths)] for label,col in [('HV',4),('Robotaxi',5)]}
    integrated = {k: sum(v) for k,v in masses.items()}
    means = {k: sum(r[3]*w for r,w in zip(rows,v))/sum(v) for k,v in masses.items()}
    bins_overlap = sum(min(a/integrated['HV'],b/integrated['Robotaxi']) for a,b in zip(masses['HV'],masses['Robotaxi']))
    r4b = list(book['figure4b'].iter_rows(min_row=2,values_only=True))
    bins4b = [r for r in r4b if isinstance(r[0], (int, float))]
    assert [r[0] for r in bins4b] == list(range(1, 16))
    totals = {name: sum(float(r[i]) for r in bins4b) for i,name in
              [(1,'observed'),(2,'CTX_predicted'),(3,'POP_SNE_predicted')]}
    declared_n = next(r[1] for r in r4b if r[0] == 'N')
    assert totals['observed'] == declared_n
    book.close()
    return {'figure4a_bins': len(rows), 'published_statistics': published,
            'rounded_density_integrals': integrated,
            'renormalized_bin_center_mean_approximation': means,
            'renormalized_binned_distance_distribution_overlap_approximation': bins_overlap,
            'figure4b_cci_bins': len(bins4b), 'figure4b_totals': totals,
            'figure4b_declared_n': declared_n,
            'figure4b_exclusion_note': next(r[0] for r in r4b if isinstance(r[0], str) and r[0].startswith('Compared')),
            'join_limit': 'No per-trip/grid/time keys linking Figure4a histogram to Figure2b; common-area standardization unavailable',
            'weighting': 'Figure4a upstream code weights trip SNE_mean by dist_km; not order-count shares'}


def main():
    started = time.perf_counter()
    inputs = {f'figure{n}': PUBLIC/'publisher'/f'41893_2026_1944_MOESM{file}_ESM.xlsx' for n,file in [(2,10),(4,12)]}
    inputs.update({name: PUBLIC/'repository/6-figure'/f'{name}.py' for name in ['fig2b','fig2c','fig4a','fig4b']})
    before = {k:sha(v) for k,v in inputs.items()}
    result = {'scope':'published aligned spatial activity; no capability or causal inference',
              'sources': {k:{'path':str(v.relative_to(ROOT)), 'sha256':before[k]} for k,v in inputs.items()},
              'figure2': audit_figure2(inputs['figure2']), 'figure4': audit_figure4(inputs['figure4'])}
    assert before == {k:sha(v) for k,v in inputs.items()}
    result['inputs_unchanged'] = True
    result['runtime_s'] = time.perf_counter()-started
    try:
        import psutil
        mem=psutil.Process().memory_info()
        result['peak_working_set_mib'] = getattr(mem,'peak_wset',mem.rss)/2**20
    except ImportError:
        import ctypes
        from ctypes import wintypes
        class Counters(ctypes.Structure):
            _fields_ = [('cb', wintypes.DWORD), ('faults', wintypes.DWORD)] + [(k, ctypes.c_size_t) for k in
                ['peak_working_set', 'working_set', 'peak_paged', 'paged', 'peak_nonpaged', 'nonpaged', 'pagefile', 'peak_pagefile']]
        counters = Counters(); counters.cb = ctypes.sizeof(counters)
        kernel = ctypes.WinDLL('kernel32'); kernel.GetCurrentProcess.restype = wintypes.HANDLE
        psapi = ctypes.WinDLL('psapi')
        psapi.GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.POINTER(Counters), wintypes.DWORD]
        ok = psapi.GetProcessMemoryInfo(kernel.GetCurrentProcess(), ctypes.byref(counters), counters.cb)
        result['peak_working_set_mib'] = counters.peak_working_set / 2**20 if ok else None
    DOC.mkdir(parents=True,exist_ok=True)
    (DOC/'summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
    print('COMPLETE',result['runtime_s'],result['peak_working_set_mib'],flush=True)


if __name__=='__main__':
    main()
