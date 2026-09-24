"""Read downloaded public inputs without executing upstream code or exporting identifiers."""
from pathlib import Path
import csv
import json
import hashlib
from collections import Counter
from datetime import datetime
from statistics import quantiles, median
import openpyxl
from acquire_wuhan_robotaxi_public import ROOT

def main():
    result = {'scope': 'read-only input inventory; not model validation', 'csv': {}, 'workbooks': {}}
    for path in ROOT.glob('repository/7-sample_data_and_code/**/*.csv'):
        with path.open(encoding='utf-8-sig', newline='') as f:
            reader = csv.DictReader(f)
            cols = reader.fieldnames
            n = 0
            ids, vehicles = set(), set()
            counts = {k: Counter() for k in ['订单状态','订单来源','weekday','hour'] if k in cols}
            missing = Counter()
            intervals = {k: [] for k in ['request_to_arrival_s','request_to_trip_start_s','request_to_accept_s']}
            errors = Counter()
            for row in reader:
                n += 1
                for col in cols:
                    if row.get(col) in ('', None): missing[col] += 1
                # No raw IDs or addresses are emitted.
                id_col = '订单号' if '订单号' in cols else 'order_id'
                v_col = '车辆id' if '车辆id' in cols else 'vehicle_id'
                if id_col in cols: ids.add(row[id_col])
                if v_col in cols: vehicles.add(row[v_col])
                for key, counter in counts.items(): counter[row.get(key)] += 1
                if '呼单时间' in cols:
                    for key, end in [('request_to_arrival_s','到达起点时间'),('request_to_trip_start_s','开始行程时间'),('request_to_accept_s','接单时间')]:
                        try:
                            start_dt = datetime.strptime(row['呼单时间'], '%Y/%m/%d %H:%M')
                            end_dt = datetime.strptime(row[end], '%Y/%m/%d %H:%M')
                            intervals[key].append((end_dt-start_dt).total_seconds())
                        except (ValueError, KeyError): errors[key] += 1
            stats = {}
            for key, values in intervals.items():
                if not values: continue
                q = quantiles(values, n=100, method='inclusive')
                stats[key] = {'n':len(values),'missing_or_parse_error':errors[key],'min':min(values),'p25':q[24],'p50':median(values),'p75':q[74],'p90':q[89],'max':max(values),'negative':sum(v<0 for v in values),'zero':sum(v==0 for v in values)}
            result['csv'][str(path.relative_to(ROOT))] = {'rows':n,'columns':cols,'unique_orders':len(ids) if ids else None,'unique_vehicles':len(vehicles) if vehicles else None,'missing':dict(missing),'categories':{k:dict(v) for k,v in counts.items()},'timing_seconds':stats}
    for path in sorted((ROOT/'publisher').glob('*.xlsx')):
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        result['workbooks'][path.name] = {s.title:{'declared_rows_including_header':s.max_row,'columns':s.max_column,'header':list(next(s.iter_rows(max_row=1,values_only=True)))} for s in book}
        if 'figure4a' in book:
            result['figure4a_published_statistics'] = dict(book['figure4a'].iter_rows(min_row=2,max_row=8,min_col=7,max_col=8,values_only=True))
        book.close()
    result['lfs_checks'] = []
    for pointer in (ROOT/'lfs_pointers').rglob('*'):
        if not pointer.is_file(): continue
        p = ROOT/'repository'/pointer.relative_to(ROOT/'lfs_pointers')
        if not p.is_file(): continue
        expected = pointer.read_text().split('oid sha256:')[1].splitlines()[0]
        with p.open('rb') as f: actual = hashlib.file_digest(f,'sha256').hexdigest()
        result['lfs_checks'].append({'path':str(p.relative_to(ROOT)),'sha256':actual,'matches':actual==expected})
    result['downloaded_bytes'] = sum(p.stat().st_size for p in ROOT.rglob('*') if p.is_file())
    (ROOT/'inspection_summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k not in ['csv','workbooks']},ensure_ascii=True,indent=2))
    for name, v in result['csv'].items():
        print(name, v['rows'], json.dumps(v['timing_seconds']))

if __name__ == '__main__':
    main()
